"""Download the official UCI archive and import historical invoices atomically."""

import argparse
import hashlib
import json
import shutil
import sqlite3
import tempfile
import zipfile
from collections import Counter
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

import requests

from .store import create_schema, decimal_text

DATA_DIR = Path(__file__).resolve().parents[2] / ".tmp" / "support-app"
DATABASE = DATA_DIR / "uci-retail.sqlite"
URL = "https://archive.ics.uci.edu/static/public/352/online%2Bretail.zip"
HEADERS = (
    "InvoiceNo",
    "StockCode",
    "Description",
    "Quantity",
    "InvoiceDate",
    "UnitPrice",
    "CustomerID",
    "Country",
)


def import_rows(rows, path: Path, provenance: dict) -> dict:
    """Keep valid identified-customer rows, preserve cancellations and exact monetary values."""
    path.parent.mkdir(parents=True, exist_ok=True)
    db = sqlite3.connect(path)
    counts = Counter()
    try:
        create_schema(db)
        if db.execute("SELECT 1 FROM metadata WHERE key='import_complete'").fetchone():
            raise ValueError("Already imported; refusing to overwrite orders or existing tickets.")
        orders = {}
        with db:
            for source_row, row in enumerate(rows, start=2):
                counts["source_rows"] += 1
                invoice, code, description, quantity, date, price, customer, country = row
                if customer is None:
                    counts["excluded_missing_customer"] += 1
                    continue
                if not isinstance(date, datetime) or not invoice or not code:
                    counts["excluded_invalid_fields"] += 1
                    continue
                try:
                    qty_decimal, money = Decimal(str(quantity)), Decimal(str(price))
                    customer_decimal = Decimal(str(customer))
                    if not all(x.is_finite() for x in (qty_decimal, money, customer_decimal)):
                        raise ValueError("Non-finite numeric field")
                    if (
                        qty_decimal != qty_decimal.to_integral_value()
                        or customer_decimal != customer_decimal.to_integral_value()
                    ):
                        raise ValueError("Expected integer quantity and customer")
                    qty, customer_id = int(qty_decimal), str(int(customer_decimal))
                    if qty == 0 or money < 0:
                        raise ValueError("Unsupported quantity/price")
                except (InvalidOperation, ValueError, OverflowError):
                    counts["excluded_invalid_fields"] += 1
                    continue
                invoice = str(invoice).strip().upper()
                record_type = (
                    "cancellation_record" if invoice.startswith("C") else "purchase_record"
                )
                amount = money * qty
                key = (customer_id, date.isoformat(), str(country or "Unknown"), record_type)
                if invoice in orders:
                    if (orders[invoice]["key"][0], orders[invoice]["key"][2:]) != (key[0], key[2:]):
                        raise ValueError(f"Inconsistent invoice ownership or country: {invoice}")
                    if orders[invoice]["key"][1] != key[1]:
                        counts["lines_with_different_invoice_time"] += 1
                        previous = orders[invoice]["key"]
                        orders[invoice]["key"] = (
                            previous[0],
                            min(previous[1], key[1]),
                            *previous[2:],
                        )
                else:
                    orders[invoice] = {"key": key, "total": Decimal(0), "count": 0}
                orders[invoice]["total"] += amount
                orders[invoice]["count"] += 1
                db.execute(
                    "INSERT INTO order_lines VALUES (?, ?, ?, ?, ?, ?, ?, ?)",
                    (
                        source_row,
                        invoice,
                        str(code),
                        str(description or "Description unavailable"),
                        qty,
                        decimal_text(money),
                        decimal_text(amount),
                        date.isoformat(),
                    ),
                )
                counts["imported_lines"] += 1
            if not orders:
                raise ValueError("No usable customer invoices found.")
            db.executemany(
                "INSERT INTO orders VALUES (?, ?, ?, ?, ?, ?, ?)",
                [
                    (invoice, *value["key"], decimal_text(value["total"]), value["count"])
                    for invoice, value in orders.items()
                ],
            )
            counts["invoices"] = len(orders)
            counts["customers"] = len({value["key"][0] for value in orders.values()})
            report = {**provenance, **dict(counts)}
            db.execute("INSERT INTO metadata VALUES ('provenance', ?)", (json.dumps(report),))
            db.execute("INSERT INTO metadata VALUES ('import_complete', '1')")
        return report
    finally:
        db.close()


def download_workbook(destination: Path) -> str:
    """Read only the expected workbook member, never extract arbitrary archive paths."""
    destination.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(dir=destination.parent) as temp:
        archive = Path(temp) / "data.zip"
        sha = hashlib.sha256()
        size = 0
        with requests.get(URL, stream=True, timeout=(15, 90)) as response:
            response.raise_for_status()
            with archive.open("wb") as out:
                for chunk in response.iter_content(1024 * 1024):
                    size += len(chunk)
                    if size > 40_000_000:
                        raise ValueError("Unexpected archive size")
                    sha.update(chunk)
                    out.write(chunk)
        with zipfile.ZipFile(archive) as bundle:
            info = bundle.getinfo("Online Retail.xlsx")
            if info.file_size > 40_000_000:
                raise ValueError("Unexpected workbook size")
            with bundle.open(info) as source, destination.open("wb") as out:
                shutil.copyfileobj(source, out)
        return sha.hexdigest()


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--xlsx", type=Path, help="Use an already downloaded official workbook")
    args = parser.parse_args()
    if DATABASE.exists():
        with sqlite3.connect(DATABASE) as db:
            if db.execute("SELECT value FROM metadata WHERE key='import_complete'").fetchone():
                print("Dataset already imported; existing tickets preserved.")
                return
    try:
        import openpyxl
    except ImportError:
        raise SystemExit(
            "Install importer requirements: uv pip install -r examples/support_app/requirements-data.txt"
        ) from None
    workbook_path = args.xlsx or DATA_DIR / "source" / "Online Retail.xlsx"
    archive_sha = None
    if not args.xlsx:
        print("Downloading official UCI Online Retail dataset…", flush=True)
        archive_sha = download_workbook(workbook_path)
    workbook_sha = hashlib.sha256(workbook_path.read_bytes()).hexdigest()
    book = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        rows = book.active.iter_rows(values_only=True)
        if tuple(next(rows)) != HEADERS:
            raise ValueError("Workbook headers do not match the official dataset.")
        print("Importing historical invoices into SQLite…", flush=True)
        report = import_rows(
            rows,
            DATABASE,
            {
                "source": "https://archive.ics.uci.edu/dataset/352/online+retail",
                "doi": "10.24432/C5BW33",
                "license": "CC BY 4.0",
                "period": "2010-12-01 to 2011-12-09",
                "workbook_sha256": workbook_sha,
                "archive_sha256": archive_sha,
            },
        )
        print(json.dumps(report, indent=2))
    finally:
        book.close()


if __name__ == "__main__":
    main()
